import io
import json
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.text import PP_ALIGN

st.set_page_config(page_title="pHoreBrain Studio", page_icon="🧠", layout="wide")

# pHoreBrain Studio — MVP 0.1
GREEN = "#B5F36A"
DARK = "#102B24"


st.markdown("""
<style>
/* Main app */
.stApp {
    background-color: #101714;
    color: #F2F7F4;
}

/* Main page text */
.stApp h1, .stApp h2, .stApp h3,
.stApp p, .stApp li, .stApp label,
.stApp [data-testid="stMarkdownContainer"] {
    color: #F2F7F4;
}

/* Sidebar */
[data-testid="stSidebar"] {
    background-color: #102B24;
}
[data-testid="stSidebar"] * {
    color: #F2F7F4 !important;
}

/* Input fields and text areas */
.stApp input,
.stApp textarea {
    background-color: #1B2721 !important;
    color: #FFFFFF !important;
    -webkit-text-fill-color: #FFFFFF !important;
    caret-color: #FFFFFF;
}
.stApp input::placeholder,
.stApp textarea::placeholder {
    color: #FFFFFF !important;
    -webkit-text-fill-color: #C0CEC6 !important;
    opacity: 1;
}

/* Dropdowns */
.stApp [data-baseweb="select"] > div {
    background-color: #1B2721;
    color: #FFFFFF;
}
.stApp [data-baseweb="select"] * {
    color: #FFFFFF;
}

/* Hero banner */
.hero {
    background-color: #102B24;
    color: #FFFFFF;
    padding: 1.5rem;
    border-radius: 16px;
}
.hero h1 {
    color: #FFFFFF !important;
    margin: 0;
}
.hero p {
    color: #D2E1D9 !important;
    margin-bottom: 0;
}

/* Buttons */
div.stButton > button,
div.stDownloadButton > button,
.stFormSubmitButton > button {
    border-radius: 10px;
}

/* Expanders */
.stApp [data-testid="stExpander"] {
    background-color: #1B2721;
}

/* Improve contrast for secondary text */
.stApp [data-testid="stCaptionContainer"] {
    color: #D2E1D9;
}
</style>
""", unsafe_allow_html=True)


if "questions" not in st.session_state:
    st.session_state.questions = []

def sample_questions():
    return [
        {
            "question": "Which organelle is known as the powerhouse of the cell?",
            "options": ["Nucleus", "Mitochondrion", "Ribosome", "Golgi apparatus"],
            "answer": 1,
            "explanation": "Mitochondria produce most of the ATP used by the cell."
        },
        {
            "question": "What is the SI unit of force?",
            "options": ["Joule", "Watt", "Newton", "Pascal"],
            "answer": 2,
            "explanation": "Force is measured in newtons (N) in the SI system."
        }
    ]

def build_docx(title, subject, chapter, questions, include_answers, include_explanations):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title.strip() or "Practice MCQ Set")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(16, 43, 36)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(" | ".join(x for x in [subject.strip(), chapter.strip()] if x)).italic = True
    doc.add_paragraph("")

    for i, q in enumerate(questions, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f"{i}. {q['question']}")
        run.bold = True
        run.font.size = Pt(11)
        for j, opt in enumerate(q["options"]):
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Inches(0.25)
            op.paragraph_format.space_after = Pt(1)
            op.add_run(f"{chr(65+j)}. {opt}")
        if include_answers:
            ans = doc.add_paragraph()
            ans.paragraph_format.left_indent = Inches(0.25)
            ar = ans.add_run(f"Answer: {chr(65+q['answer'])}. {q['options'][q['answer']]}")
            ar.bold = True
            ar.font.color.rgb = RGBColor(35, 115, 68)
        if include_explanations and q.get("explanation", "").strip():
            ep = doc.add_paragraph()
            ep.paragraph_format.left_indent = Inches(0.25)
            ep.add_run("Explanation: ").bold = True
            ep.add_run(q["explanation"])
        doc.add_paragraph("")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()

def build_pptx(title, subject, chapter, questions, include_answers):
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = PptRGBColor(16, 43, 36)
    shape = slide.shapes.add_textbox(PptInches(0.8), PptInches(1.55), PptInches(11.7), PptInches(2.0))
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title.strip() or "Practice MCQ Set"
    p.font.size = PptPt(34)
    p.font.bold = True
    p.font.color.rgb = PptRGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    sub = slide.shapes.add_textbox(PptInches(1.0), PptInches(4.0), PptInches(11.3), PptInches(0.8))
    sp = sub.text_frame.paragraphs[0]
    sp.text = " • ".join(x for x in [subject.strip(), chapter.strip()] if x)
    sp.font.size = PptPt(20)
    sp.font.color.rgb = PptRGBColor(181, 243, 106)
    sp.alignment = PP_ALIGN.CENTER

    for i, q in enumerate(questions, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PptRGBColor(248, 250, 248)
        top = slide.shapes.add_shape(1, 0, 0, prs.slide_width, PptInches(0.22))
        top.fill.solid()
        top.fill.fore_color.rgb = PptRGBColor(181, 243, 106)
        top.line.fill.background()

        qbox = slide.shapes.add_textbox(PptInches(0.7), PptInches(0.65), PptInches(12), PptInches(1.6))
        qtf = qbox.text_frame
        qtf.word_wrap = True
        qp = qtf.paragraphs[0]
        qp.text = f"Q{i}. {q['question']}"
        qp.font.size = PptPt(26)
        qp.font.bold = True
        qp.font.color.rgb = PptRGBColor(16, 43, 36)

        obox = slide.shapes.add_textbox(PptInches(1.0), PptInches(2.5), PptInches(11.5), PptInches(3.4))
        otf = obox.text_frame
        otf.word_wrap = True
        otf.clear()
        for j, opt in enumerate(q["options"]):
            p = otf.paragraphs[0] if j == 0 else otf.add_paragraph()
            p.text = f"{chr(65+j)}.  {opt}"
            p.font.size = PptPt(22)
            p.font.color.rgb = PptRGBColor(40, 55, 48)
            p.space_after = PptPt(16)

        if include_answers:
            abox = slide.shapes.add_textbox(PptInches(0.9), PptInches(6.35), PptInches(11.5), PptInches(0.5))
            ap = abox.text_frame.paragraphs[0]
            ap.text = f"Answer: {chr(65+q['answer'])}. {q['options'][q['answer']]}"
            ap.font.size = PptPt(17)
            ap.font.bold = True
            ap.font.color.rgb = PptRGBColor(35, 115, 68)

        footer = slide.shapes.add_textbox(PptInches(11.3), PptInches(7.05), PptInches(1.3), PptInches(0.25))
        fp = footer.text_frame.paragraphs[0]
        fp.text = f"{i}/{len(questions)}"
        fp.font.size = PptPt(10)
        fp.font.color.rgb = PptRGBColor(100, 115, 105)
        fp.alignment = PP_ALIGN.RIGHT

    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()

with st.sidebar:
    st.markdown("## 🧠 pHoreBrain Studio")
    st.caption("Question Bank & Document Builder")
    page = st.radio("Workspace", ["Question Editor", "Export Center", "About MVP"])
    st.markdown("---")
    st.caption("MVP 0.1 • Local-first prototype")

st.markdown(f"""
<div class="hero">
  <h1>pHoreBrain Studio</h1>
  <p>One question bank. Multiple study resources.</p>
</div>
""", unsafe_allow_html=True)

if page == "Question Editor":
    st.subheader("1. Set up your resource")
    c1, c2, c3 = st.columns(3)
    with c1:
        title = st.text_input("Resource title", "Biology & Physics Practice Set", key="resource_title")
    with c2:
        subject = st.selectbox("Subject", ["Biology", "Chemistry", "Physics", "English", "General Knowledge", "Mixed"])
    with c3:
        chapter = st.text_input("Chapter / topic", "Cell Biology / Mechanics")

    st.subheader("2. Add a multiple-choice question")
    with st.form("question_form", clear_on_submit=True):
        question = st.text_area("Question text", placeholder="Type your question here...")
        col1, col2 = st.columns(2)
        with col1:
            a = st.text_input("Option A")
            b = st.text_input("Option B")
        with col2:
            c = st.text_input("Option C")
            d = st.text_input("Option D")
        answer_label = st.selectbox("Correct answer", ["A", "B", "C", "D"])
        explanation = st.text_area("Explanation (optional)")
        submitted = st.form_submit_button("＋ Add question", use_container_width=True)
        if submitted:
            opts = [a.strip(), b.strip(), c.strip(), d.strip()]
            if not question.strip():
                st.error("Please enter the question text.")
            elif not all(opts):
                st.error("Please fill in all four options.")
            else:
                st.session_state.questions.append({
                    "question": question.strip(),
                    "options": opts,
                    "answer": ["A", "B", "C", "D"].index(answer_label),
                    "explanation": explanation.strip()
                })
                st.success("Question added to this session.")

    b1, b2 = st.columns(2)
    with b1:
        if st.button("Load sample questions", use_container_width=True):
            st.session_state.questions = sample_questions()
            st.rerun()
    with b2:
        if st.button("Clear all questions", use_container_width=True):
            st.session_state.questions = []
            st.rerun()

    st.subheader(f"3. Question bank preview ({len(st.session_state.questions)} questions)")
    if not st.session_state.questions:
        st.info("Add a question above or load the sample questions to test the builder.")
    else:
        for i, q in enumerate(st.session_state.questions):
            with st.expander(f"Q{i+1}. {q['question']}", expanded=(i == len(st.session_state.questions)-1)):
                st.write("  |  ".join(f"{chr(65+j)}. {opt}" for j, opt in enumerate(q["options"])))
                st.caption(f"Correct answer: {chr(65+q['answer'])}")
                if q.get("explanation"):
                    st.write(q["explanation"])
                if st.button("Delete this question", key=f"del_{i}"):
                    st.session_state.questions.pop(i)
                    st.rerun()

    if st.session_state.questions:
        st.download_button(
            "Download question bank JSON",
            data=json.dumps({"title": title, "subject": subject, "chapter": chapter, "questions": st.session_state.questions}, ensure_ascii=False, indent=2),
            file_name="phorebrain_question_bank.json",
            mime="application/json",
            use_container_width=True
        )

elif page == "Export Center":
    st.subheader("Export your resource")
    title = st.text_input("Resource title", st.session_state.get("resource_title", "Practice MCQ Set"), key="export_title")
    subject = st.selectbox("Subject", ["Biology", "Chemistry", "Physics", "English", "General Knowledge", "Mixed"], key="export_subject")
    chapter = st.text_input("Chapter / topic", "Practice Set", key="export_chapter")
    if not st.session_state.questions:
        st.warning("Your question bank is empty. Go to Question Editor and add questions first.")
    else:
        st.write(f"Ready to export **{len(st.session_state.questions)} questions**.")
        include_answers = st.checkbox("Include correct answers in output", value=False)
        include_explanations = st.checkbox("Include explanations in DOCX", value=True)
        col1, col2 = st.columns(2)
        with col1:
            docx_bytes = build_docx(title, subject, chapter, st.session_state.questions, include_answers, include_explanations)
            st.download_button(
                "⬇ Download editable DOCX",
                data=docx_bytes,
                file_name="phorebrain_mcq_set.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )
        with col2:
            pptx_bytes = build_pptx(title, subject, chapter, st.session_state.questions, include_answers)
            st.download_button(
                "⬇ Download editable PPTX",
                data=pptx_bytes,
                file_name="phorebrain_slides.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                use_container_width=True
            )
        st.caption("Tip: keep answers turned off when making student-facing slides. Export a separate answer-key version for teachers.")

elif page == "About MVP":
    st.subheader("What this first version can do")
    st.markdown("""
- Create MCQs with four options and one correct answer.
- Add an optional explanation.
- Preview and delete questions.
- Export the question bank as JSON.
- Generate editable Word (`.docx`) question sheets.
- Generate editable PowerPoint (`.pptx`) slides.
- Choose whether answers appear in the exported materials.

**Current limitations**
- Questions are stored in the current browser session; refreshes or sleeping sessions may clear them.
- No user login, permanent database, PDF export, bulk DOCX import, image upload, or AI question generation yet.
- Review generated documents and answer keys before distributing them to students.
""")
    st.markdown("### Planned next")
    st.write("Persistent question bank → DOCX/PDF templates → slide themes → online quiz mode → user accounts.")
